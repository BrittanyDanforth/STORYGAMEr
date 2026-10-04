from pathlib import Path
import re
from collections import Counter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / 'Hacker_Game_Master_Plan.docx'
content = (ROOT/'master_content.md').read_text()
catalog = (ROOT/'feature_catalog.md').read_text()
entries = re.findall(r'^F(\d{3}) \[([SCL])\] (.+)$', catalog, flags=re.M)
assert [int(e[0]) for e in entries] == list(range(1,321))
assert Counter(e[1] for e in entries) == {'S':40,'C':217,'L':63}
catalog = re.sub(r'^## (\d{2})\. (.+)$', lambda m: '### '+str(int(m[1]))+' '+re.sub(r'[^A-Za-z0-9 ]','',m[2]), catalog, flags=re.M)
content = content.replace('{{FEATURE_CATALOG}}',catalog)
assert '{{' not in content
chapters = re.findall(r'^## (.+)$',content,flags=re.M)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.7)
sec.left_margin, sec.right_margin = Inches(.78), Inches(.78)
sec.header_distance, sec.footer_distance = Inches(.30), Inches(.32)
styles = doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st = styles[name]
    st.font.name='Liberation Sans'
    st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().append(OxmlElement('w:rFonts')) if st._element.rPr is None else None
for border in list(styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
normal=styles['Normal']
normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
styles['Title'].font.size=Pt(25)
styles['Title'].font.bold=True
styles['Title'].paragraph_format.space_after=Pt(7)
styles['Subtitle'].font.size=Pt(12)
styles['Subtitle'].paragraph_format.space_after=Pt(7)
for name,size,before,after in [('Heading 1',18,17,8),('Heading 2',13,12,5),('Heading 3',11.5,10,4)]:
    st=styles[name]
    st.font.size=Pt(size)
    st.font.bold=True
    st.paragraph_format.space_before=Pt(before)
    st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True
    st.paragraph_format.keep_together=True
for name in ['Feature','Navigation','Source']:
    st=styles.add_style(name,1)
    st.base_style=styles['Normal']
styles['Feature'].font.size=Pt(10.5)
styles['Feature'].paragraph_format.line_spacing=1.1
styles['Feature'].paragraph_format.space_after=Pt(5)
styles['Feature'].paragraph_format.keep_together=True
styles['Navigation'].font.size=Pt(10.5)
styles['Navigation'].paragraph_format.space_after=Pt(3)
styles['Navigation'].paragraph_format.line_spacing=1.0
styles['Source'].font.size=Pt(10)

def bookmark(p,name,index):
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(index)); start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(index))
    p._p.insert(0,start); p._p.append(end)

def link(p,label,target,internal=False):
    h=OxmlElement('w:hyperlink')
    if internal: h.set(qn('w:anchor'),target)
    else: h.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'000000' if internal else '215478');props.append(color)
    r.append(props);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def field(p,instruction):
    run=p.add_run(); begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=instruction
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate')
    text=OxmlElement('w:t');text.text='1'
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    for x in [begin,ins,sep,text,end]:run._r.append(x)

def prose(text):
    p=doc.add_paragraph()
    for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        run=p.add_run(part)
        if i%2:run.bold=True
    return p

def table(block):
    rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in block.splitlines()]
    rows=[r for r in rows if not all(re.fullmatch(r'[- :]+',c) for c in r)]
    count=len(rows[0]); t=doc.add_table(rows=1,cols=count)
    t.autofit=False
    widths=[1.25,2.77,2.92] if count==3 else [2.00,4.94]
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    props.append(borders)
    for ridx,row in enumerate(rows):
        tr=t.rows[0] if ridx==0 else t.add_row()
        trPr=tr._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if ridx==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for cidx,txt in enumerate(row):
            c=tr.cells[cidx];c.width=Inches(widths[cidx]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'95');el.set(qn('w:type'),'dxa');margins.append(el)
            tcPr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'203C50' if ridx==0 else ('F3F5F6' if ridx%2==0 else 'FFFFFF'));tcPr.append(shade)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
            run=p.add_run(txt);run.font.size=Pt(10.5);run.font.bold=(ridx==0);run.font.color.rgb=RGBColor.from_string('FFFFFF' if ridx==0 else '000000')
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

hp=sec.header.paragraphs[0]
hp.text=''
hp.paragraph_format.space_after=Pt(0)
fp=sec.footer.paragraphs[0]
fp.paragraph_format.tab_stops.clear_all()
fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.94),WD_TAB_ALIGNMENT.RIGHT)
fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
field(fp,'PAGE')
for run in fp.runs:run.font.size=Pt(8)

doc.add_paragraph('Open World Hacker Game Master Plan','Title')
doc.add_paragraph('Game vision systems and production roadmap','Subtitle')
p=doc.add_paragraph('Prepared for Elijah  •  Expanded planning version 0.9  •  4 October 2026')
p.runs[0].font.size=Pt(9)
doc.add_paragraph('An underworld-first, first-person 3D hacker crime and fraud simulation. Start broke and behind on rent, learn through paid work, and choose how far to go. Hacking drives income and progression. This plan connects 320 prioritized features across active NPC computers, criminal economies, hands-on PC building, hacker crews, investigations, courts, jail, a viable white-hat career, and a concrete Blender room and character build plan.')
p=doc.add_paragraph()
p.add_run('Starting direction  ').bold=True
p.add_run('Unity 6 LTS, single-player Windows, complete offline play, and optional small-model AI dialogue and bounded planning. Performance targets and content counts remain proposals until tested.')
doc.add_heading('Contents',level=2)
for i,ch in enumerate(chapters,1):
    p=doc.add_paragraph(style='Navigation')
    link(p,ch,'chapter_'+str(i),internal=True)
doc.add_page_break()

start=content.index('## 1 The game we are building')
body=content[start:]
chapter_idx=0
for block in re.split(r'\n\s*\n',body):
    block=block.strip()
    if not block:continue
    if block.startswith('### '):
        heading=re.sub(r'[^A-Za-z0-9 ]','',block[4:])
        doc.add_heading(heading,level=2)
    elif block.startswith('## '):
        chapter_idx+=1
        heading=block[3:]
        p=doc.add_heading(heading,level=1)
        if heading.startswith('25 ') or heading.startswith('27 '):p.paragraph_format.page_break_before=True
        bookmark(p,'chapter_'+str(chapter_idx),chapter_idx)
    elif block.startswith('|'):
        table(block)
    elif re.match(r'^F\d{3} ',block):
        p=doc.add_paragraph(style='Feature')
        prefix,desc=block.split('] ',1)
        p.add_run(prefix+']  ').bold=True
        p.add_run(desc)
    elif block.startswith('[R'):
        lines=block.splitlines()
        p=doc.add_paragraph(lines[0],style='Source')
        if len(lines)>1:
            p.add_run(' ')
            link(p,'Open official source',lines[1])
        p.paragraph_format.keep_together=True
    else:
        prose(block.replace('\n',' '))

doc.core_properties.title='Open World Hacker Game Master Plan'
doc.core_properties.subject='Game design and production plan'
doc.core_properties.author=''
doc.core_properties.keywords='hacker game, first person, simulation, Unity, game design'
doc.core_properties.comments=''
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.save(OUT)
print({'file':str(OUT),'words':len(content.split()),'features':len(entries),'priorities':dict(Counter(e[1] for e in entries)),'chapters':len(chapters)})
